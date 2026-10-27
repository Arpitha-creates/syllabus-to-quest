"""File ingestion: turn an uploaded syllabus or notes file into clean text.

Supported: .pdf (text layer), .docx, .txt, .md. Scanned PDFs have no text
layer - the caller gets a clear ValueError instead of an empty quest board.
"""

from __future__ import annotations

import io
from pathlib import Path

from .nlp import clean_text

SUPPORTED = (".pdf", ".docx", ".txt", ".md")


def kind_of(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(
            f"'{suffix or filename}' is not supported. Upload one of: "
            + ", ".join(SUPPORTED)
        )
    return suffix.lstrip(".")


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """Return (clean_text, kind) for an uploaded file."""
    kind = kind_of(filename)
    if kind in ("txt", "md"):
        text = _decode(data)
    elif kind == "pdf":
        text = _from_pdf(data)
    else:
        text = _from_docx(data)
    text = clean_text(text)
    if len(text) < 60:
        raise ValueError(
            "Almost no text came out of that file. If it is a scanned PDF, "
            "run OCR on it first (or paste the syllabus as text)."
        )
    return text, kind


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency hint
        raise ValueError("PDF support needs pypdf: pip install pypdf") from exc

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:                      # a broken page must not kill the upload
            pages.append("")
    return "\n".join(pages)


def _from_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dependency hint
        raise ValueError("DOCX support needs python-docx: pip install python-docx") from exc

    document = Document(io.BytesIO(data))
    chunks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)
